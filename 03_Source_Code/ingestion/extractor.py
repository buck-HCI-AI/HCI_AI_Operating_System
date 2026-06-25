#!/usr/bin/env python3
"""
Text extractor for the Knowledge Ingestion Engine.
Supports: PDF, DOCX, XLSX, TXT, MD, CSV
Returns plain text string (empty string on failure or unsupported format).
"""
import os
from typing import Optional

ALLOWED_EXTENSIONS = {
    ".pdf", ".docx", ".xlsx", ".xls",
    ".txt", ".md", ".csv",
    ".jpg", ".jpeg", ".png", ".tiff",  # accepted but no text extraction yet
}

MAX_FILE_SIZE_MB = 250


def allowed(path: str) -> bool:
    ext = os.path.splitext(path)[1].lower()
    return ext in ALLOWED_EXTENSIONS


def within_size_limit(path: str) -> bool:
    try:
        return os.path.getsize(path) <= MAX_FILE_SIZE_MB * 1024 * 1024
    except OSError:
        return False


def extract(path: str) -> str:
    """Extract text from file at path. Returns empty string on failure."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return _pdf(path)
        if ext == ".docx":
            return _docx(path)
        if ext in (".xlsx", ".xls"):
            return _xlsx(path)
        if ext in (".txt", ".md", ".csv"):
            return _plaintext(path)
    except Exception:
        pass
    return ""


def extract_bytes(data: bytes, filename: str) -> str:
    """Extract text from raw bytes. Writes to a temp file and delegates."""
    import tempfile
    ext = os.path.splitext(filename)[1].lower()
    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        return extract(tmp_path)
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


def _pdf(path: str) -> str:
    import pdfplumber
    pages = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                pages.append(text)
    return "\n".join(pages)


def _docx(path: str) -> str:
    from docx import Document
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        rows.append(f"[Sheet: {sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append("\t".join(cells))
    wb.close()
    return "\n".join(rows)


def _plaintext(path: str) -> str:
    with open(path, encoding="utf-8", errors="ignore") as f:
        return f.read()
